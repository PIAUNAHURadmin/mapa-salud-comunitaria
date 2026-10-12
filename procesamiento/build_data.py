#!/usr/bin/env python3
"""
Pipeline de procesamiento de datos — Panel de Salud Comunitaria (Región Sanitaria VII).

Lee los archivos relevados por el equipo de Salud Comunitaria y produce la base de
datos normalizada y georreferenciada que consume el panel web.

Entradas (carpeta `drive-download-...`):
  * DATOS VINCULADOS A LA TABLA/**/*.xlsx  -> efectores de salud (1º/2º/3º nivel) y nutrición
  * REPOSITORIO DE ENLACES GEOREF.xlsx      -> taxonomía categoría/subcategoría/municipio

Salidas (carpeta `data/`):
  * municipios.json   -> 10 municipios de la RS VII con coordenadas de referencia
  * efectores.json    -> efectores normalizados + georreferenciados
  * taxonomia.json    -> categorías y subcategorías
  * resumen.json      -> conteos y métricas
"""
import json, os, re, time, glob, urllib.parse, urllib.request, sys

ROOT = os.path.dirname(os.path.abspath(__file__))
# Carpeta de datos crudos: `RawData` (antes se llamaba `drive-download-...`).
_CANDIDATAS = ["RawData", "drive-download-20260929T120752Z-1-001"]
SOURCE = next((os.path.abspath(os.path.join(ROOT, "..", "..", c)) for c in _CANDIDATAS
               if os.path.isdir(os.path.join(ROOT, "..", "..", c))),
              os.path.abspath(os.path.join(ROOT, "..", "..", "RawData")))
DATA = os.path.join(ROOT, "..", "data")
CACHE = os.path.join(ROOT, "geocode_cache.json")

os.makedirs(DATA, exist_ok=True)

# ---------------------------------------------------------------------------
# 1) Nombres canónicos de municipios (Región Sanitaria VII)
# ---------------------------------------------------------------------------
MUNICIPIOS_CANON = {
    "HURLINGHAM": "Hurlingham",
    "ITUZAINGO": "Ituzaingó",
    "ITUZAINGÓ": "Ituzaingó",
    "MORON": "Morón",
    "MORÓN": "Morón",
    "MERLO": "Merlo",
    "MORENO": "Moreno",
    "TRES DE FEBRERO": "Tres de Febrero",
    "GENERAL RODRIGUEZ": "General Rodríguez",
    "GENERAL RODRÍGUEZ": "General Rodríguez",
    "GENERAL LAS HERAS": "General Las Heras",
    "MARCOS PAZ": "Marcos Paz",
    "LUJAN": "Luján",
    "LUJÁN": "Luján",
}

def canon_mun(raw):
    if not raw:
        return ""
    r = raw.strip().upper()
    for k, v in MUNICIPIOS_CANON.items():
        if r == k or r.startswith(k):
            return v
    return raw.strip().title()

# Caja aproximada que contiene los 10 municipios de la Región Sanitaria VII.
# Coordenadas fuera de esta caja se consideran geocodificación errónea
# (p. ej. Nominatim devolvió una calle homónima en otro partido o en CABA).
BBOX_RS7 = {"lat_min": -35.25, "lat_max": -34.35, "lon_min": -59.50, "lon_max": -58.49}

# Coordenadas de referencia fijadas a mano cuando Nominatim resuelve mal el municipio.
# "Tres de Febrero, Buenos Aires" devuelve la Av. 3 de Febrero en Palermo (CABA);
# se usa Caseros, cabecera del partido.
MUNICIPIO_COORD_MANUAL = {
    "Tres de Febrero": {"lat": -34.6069, "lon": -58.5636,
                        "display_name": "Caseros, Partido de Tres de Febrero, Buenos Aires, Argentina (fijado a mano)"},
}

def dentro_rs7(lat, lon):
    return (BBOX_RS7["lat_min"] <= lat <= BBOX_RS7["lat_max"]
            and BBOX_RS7["lon_min"] <= lon <= BBOX_RS7["lon_max"])

# ---------------------------------------------------------------------------
# 2) Geocodificación (Nominatim / OpenStreetMap) con caché y rate-limit
# ---------------------------------------------------------------------------
def load_cache():
    if os.path.exists(CACHE):
        return json.load(open(CACHE, encoding="utf-8"))
    return {}

def save_cache(c):
    json.dump(c, open(CACHE, "w", encoding="utf-8"), ensure_ascii=False, indent=1)

cache = load_cache()

def geocode(query, country="Argentina", sleep=1.1):
    key = query.strip().lower()
    if key in cache:
        return cache[key]
    params = urllib.parse.urlencode({"q": query, "format": "json", "limit": 1, "countrycodes": "ar"})
    url = "https://nominatim.openstreetmap.org/search?" + params
    req = urllib.request.Request(url, headers={"User-Agent": "salud-comunitaria-demo/1.0"})
    try:
        with urllib.request.urlopen(req, timeout=15) as r:
            data = json.loads(r.read().decode())
        if data:
            res = {"lat": float(data[0]["lat"]), "lon": float(data[0]["lon"]),
                   "display_name": data[0].get("display_name", "")}
        else:
            res = None
    except Exception as e:
        # Los errores de red NO se guardan en caché, para reintentar en la próxima corrida.
        print("    [geocode] error de red:", e)
        return None
    cache[key] = res
    save_cache(cache)
    time.sleep(sleep)
    return res

# ---------------------------------------------------------------------------
# 3) Extracción de xlsx
# ---------------------------------------------------------------------------
import openpyxl

def extract_xlsx(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    meta, data, header_idx = {}, [], None
    for i, row in enumerate(rows):
        cells = [str(c).strip() if c is not None else '' for c in row]
        if cells and cells[0] in ('Establecimiento', 'Nombre', 'NOMBRE', 'ESTABLECIMIENTO'):
            header_idx = i
            break
        joined = " ".join(cells)
        m = re.match(r'^Partido:\s*(.*)$', joined, re.I)
        if m: meta['partido'] = m.group(1).strip()
        m = re.match(r'^(Tipo de establecimiento|Especialidad):\s*(.*)$', joined, re.I)
        if m: meta['tipo'] = m.group(2).strip()
        m = re.match(r'^Resultados obtenidos:\s*(\d+)$', joined, re.I)
        if m: meta['resultados'] = int(m.group(1))
    if header_idx is None:
        return meta, []
    headers = [str(c).strip() if c is not None else '' for c in rows[header_idx]]
    for row in rows[header_idx+1:]:
        vals = [str(c).strip() if c is not None else '' for c in row]
        if not any(vals): continue
        rec = dict(zip(headers, vals))
        if not rec.get('Establecimiento') and not rec.get('Nombre'):
            continue
        data.append(rec)
    return meta, data

def classify(path):
    p = path.replace('\\', '/')
    cat = sub = nivel = None
    if 'EFECTORES DE SALUD PUBLICA' in p:
        cat = 'EFECTORES DE SALUD PÚBLICA'
        if 'PRIMER NIVEL' in p: nivel = 'Primer Nivel'
        elif 'SEGUNDO NIVEL' in p: nivel = 'Segundo Nivel'
        elif 'TERCER NIVEL' in p: nivel = 'Tercer Nivel'
        sub = nivel
    elif 'NUTRICIÓN Y ALIMENTACIÓN' in p or 'NUTRICION' in p:
        cat = 'NUTRICIÓN Y ALIMENTACIÓN'
        sub = 'Consultorios de nutrición'
    return cat, sub, nivel

def _norm(t):
    """Mayúsculas sin acentos ni guiones bajos, para comparar nombres."""
    import unicodedata
    t = unicodedata.normalize("NFKD", t or "").encode("ascii", "ignore").decode()
    return t.upper().replace("_", " ").strip()

def partido_coincide_con_archivo(path, partido_header):
    """El nombre del archivo debe contener el partido de la cabecera.
    Detecta archivos copiados/mal nombrados (p. ej. el de Tres de Febrero
    que en realidad trae los datos de Morón)."""
    nombre = _norm(os.path.basename(path))
    return _norm(partido_header) in nombre

def extract_efectores():
    files = glob.glob(os.path.join(SOURCE, "DATOS VINCULADOS A LA TABLA", "**", "*.xlsx"), recursive=True)
    recs = []
    for f in sorted(files):
        cat, sub, nivel = classify(f)
        meta, data = extract_xlsx(f)
        if meta.get('partido') and not partido_coincide_con_archivo(f, meta['partido']):
            print(f"  [AVISO] se omite {os.path.relpath(f, SOURCE)}: "
                  f"la cabecera dice 'Partido: {meta['partido']}' y no coincide con el nombre del archivo")
            continue
        partido = canon_mun(meta.get('partido', ''))
        for rec in data:
            recs.append({
                'nombre': (rec.get('Establecimiento') or rec.get('Nombre') or '').strip(),
                'domicilio': (rec.get('Domicilio') or '').strip(),
                'telefono': (rec.get('Tel. Contacto') or rec.get('Tel') or '').strip(),
                'sitio_web': (rec.get('Sitio web') or '').strip(),
                'partido': partido,
                'categoria': cat,
                'subcategoria': sub,
                'nivel': nivel,
            })
    # deduplicar por nombre+domicilio
    seen, dedup = set(), []
    for r in recs:
        k = (r['nombre'].upper(), r['domicilio'].upper())
        if k in seen: continue
        seen.add(k)
        dedup.append(r)
    return dedup

# ---------------------------------------------------------------------------
# 3b) Unificación: un registro por establecimiento con la lista de sus servicios
# ---------------------------------------------------------------------------
def _limpio(t):
    return re.sub(r"\s+", " ", t or "").strip()

def _clave_nombre(t):
    """Nombre normalizado para comparar: sin tildes, sin signos, 'N° 5' -> '5'."""
    t = _norm(t)
    t = re.sub(r"[^A-Z0-9 ]", " ", t)
    t = re.sub(r"\bN\s*O?\s*(?=\d)", "", t)
    return re.sub(r"\s+", " ", t).strip()

def unificar(efectores):
    """Junta los registros del mismo establecimiento (mismo partido y nombre).
    Ej.: un CAPS que figura en 'Primer Nivel' y también en 'Consultorios de nutrición'
    (en los Excel su domicilio difería solo por espacios dobles)."""
    grupos = {}
    for r in efectores:
        grupos.setdefault((r['partido'], _clave_nombre(r['nombre'])), []).append(r)
    out = []
    for regs in grupos.values():
        # registro principal: el que tiene nivel de atención (viene de los listados de efectores)
        base = dict(next((r for r in regs if r.get('nivel')), regs[0]))
        # domicilios originales de todos los registros unidos (claves de la caché de geocodificación);
        # se prueba primero el del registro principal y después los demás.
        base['_doms_geo'] = [base['domicilio']] + [r['domicilio'] for r in regs if r['domicilio'] != base['domicilio']]
        base['nombre'] = _limpio(base['nombre'])
        base['domicilio'] = _limpio(base['domicilio'])
        base['telefono'] = _limpio(base['telefono'])
        # completar datos vacíos con los de los otros registros del mismo establecimiento
        for campo in ('telefono', 'sitio_web'):
            if base.get(campo) in (None, '', '-'):
                base[campo] = next((_limpio(r[campo]) for r in regs if r.get(campo) not in (None, '', '-')), '-')
        base['servicios'] = []
        for r in regs:
            sv = {'categoria': r['categoria'], 'subcategoria': r['subcategoria']}
            if sv not in base['servicios']:
                base['servicios'].append(sv)
        base['fuentes'] = sorted({r.get('_fuente', 'listados_xlsx') for r in regs})
        base.pop('_fuente', None)
        out.append(base)
    return out

def caps_tres_de_febrero(efectores):
    """Si no hay 1er nivel de Tres de Febrero (su Excel es copia del de Morón y se omite),
    se usan los CAPS de Tres de Febrero del listado provincial de CAPS (RawData/CAPS)."""
    if any(r['partido'] == 'Tres de Febrero' and r['nivel'] == 'Primer Nivel' for r in efectores):
        return []
    path = os.path.join(ROOT, "caps_tres_de_febrero.json")
    if not os.path.exists(path):
        print("  [AVISO] falta caps_tres_de_febrero.json: correr antes `python extraer_pdfs.py`")
        return []
    out = [{'nombre': c['nombre'], 'domicilio': c['domicilio'], 'telefono': c.get('telefono', '-'),
            'sitio_web': '-', 'partido': 'Tres de Febrero', 'categoria': 'EFECTORES DE SALUD PÚBLICA',
            'subcategoria': 'Primer Nivel', 'nivel': 'Primer Nivel', '_fuente': 'pdf_caps_provincial'}
           for c in json.load(open(path, encoding="utf-8"))]
    print(f"  + {len(out)} CAPS de Tres de Febrero desde RawData/CAPS (listado provincial)")
    return out

_STOP_NOMBRE = {"HOSPITAL", "CENTRO", "DE", "LA", "EL", "LOS", "LAS", "DEL", "Y", "SALUD", "GENERAL",
                "ZONAL", "AGUDOS", "UNIDAD", "SANITARIA", "MUNICIPAL", "DR", "PROFESOR", "DOCTOR",
                "INTERZONAL", "DESCENTRALIZADO", "CAPS", "ATENCION", "PRIMARIA", "BARRIO", "BO", "SALA"}
_STOP_DOM = {"Y", "E", "ENTRE", "ESQUINA", "ESQ", "NRO", "N", "S", "BO", "BARRIO", "AV", "AVENIDA", "DE", "LA", "EL"}

def _tokens_nombre(t, partido):
    return {w for w in _clave_nombre(t).split() if w not in _STOP_NOMBRE} - set(_clave_nombre(partido).split())

def _tokens_calle(dom):
    calle = re.split(r"\s-\s", dom or "")[0]          # "PALERMO 555 ENTRE ... - BARRIO ..., PARTIDO (CP)"
    return {w for w in _clave_nombre(calle).split() if w not in _STOP_DOM and len(w) > 1}

def buscar_establecimiento(fila, candidatos):
    """Devuelve (establecimiento, criterio) o (None, None). Criterios, en orden:
    1. nombre igual (normalizado)  2. nombre parecido (>= 0.85)
    3. mismas palabras clave en el nombre (ej. 'Centro de Salud 2 de Junio' = 'UNIDAD SANITARIA 2 DE JUNIO DE HURLINGHAM')
    4. misma calle/número (ej. 'Centro de Salud Barrio Mitre, Palermo N°555' = 'UNIDAD SANITARIA BARTOLOME MITRE, PALERMO 555')
    Los criterios 3 y 4 solo se aceptan si hay UN único candidato."""
    import difflib
    k = _clave_nombre(fila['nombre'])
    e = next((c for c in candidatos if _clave_nombre(c['nombre']) == k), None)
    if e: return e, "nombre"
    if candidatos:
        mejor = max(candidatos, key=lambda c: difflib.SequenceMatcher(None, _clave_nombre(c['nombre']), k).ratio())
        if difflib.SequenceMatcher(None, _clave_nombre(mejor['nombre']), k).ratio() >= 0.85:
            return mejor, "nombre parecido"
    tn = _tokens_nombre(fila['nombre'], fila['partido'])
    if tn:
        iguales = [c for c in candidatos if _tokens_nombre(c['nombre'], c['partido']) == tn]
        if len(iguales) == 1: return iguales[0], "palabras clave del nombre"
    tc = _tokens_calle(fila.get('domicilio', ''))
    if len(tc) >= 2:
        def cubre(c):   # qué parte de la calle de la fuente aparece en el domicilio del establecimiento
            o = _tokens_calle(c['domicilio']); return len(tc & o) / len(tc)
        buenos = [c for c in candidatos if cubre(c) >= 0.8]
        if len(buenos) == 1: return buenos[0], "domicilio"
    return None, None

def asignar_servicios_pdf(efectores):
    """Suma a cada establecimiento los servicios de los PDF provinciales
    (ver extraer_pdfs.py). NO agrega establecimientos nuevos: lo que no coincide
    se guarda en servicios_sin_match.json para revisión manual."""
    import difflib
    path = os.path.join(ROOT, "servicios_pdf.json")
    if not os.path.exists(path):
        print("  [AVISO] falta servicios_pdf.json: correr antes `python extraer_pdfs.py`")
        return
    filas = json.load(open(path, encoding="utf-8"))
    por_partido = {}
    for e in efectores:
        por_partido.setdefault(e['partido'], []).append(e)
    sin_match, n_ok, criterios = [], 0, []
    for f in filas:
        e, criterio = buscar_establecimiento(f, por_partido.get(f['partido'], []))
        if e is None:
            sin_match.append(f); continue
        if criterio != "nombre":
            criterios.append({"fuente": f['nombre'], "establecimiento": e['nombre'],
                              "partido": f['partido'], "criterio": criterio, "archivo": f['archivo']})
        sv = {'categoria': f['categoria'], 'subcategoria': f['subcategoria']}
        if sv not in e['servicios']:
            e['servicios'].append(sv)
        fuente = 'transcripcion_imagen' if f['archivo'].endswith(')') else 'pdf_provinciales'
        if fuente not in e['fuentes']:
            e['fuentes'].append(fuente)
        n_ok += 1
    json.dump(sin_match, open(os.path.join(ROOT, "servicios_sin_match.json"), "w", encoding="utf-8"),
              ensure_ascii=False, indent=1)
    # coincidencias que no fueron por nombre exacto: para revisión humana
    json.dump(criterios, open(os.path.join(ROOT, "servicios_coincidencias_a_revisar.json"), "w", encoding="utf-8"),
              ensure_ascii=False, indent=1)
    print(f"  servicios de PDF asignados: {n_ok}/{len(filas)} (sin coincidencia: {len(sin_match)})")

# ---------------------------------------------------------------------------
# 4) Limpieza de domicilio para geocodificación
# ---------------------------------------------------------------------------
def clean_address(domicilio, partido):
    d = domicilio
    # quitar el sufijo "- MUNICIPIO, MUNICIPIO (CP)"
    d = re.sub(r'\s*-\s*[^,]+,\s*[^()]*\(\d+\)\s*$', '', d)
    d = re.sub(r'\(\d{4}\)\s*$', '', d)
    d = d.replace(' - ', ', ').replace('  ', ' ').strip(' ,-')
    if not d:
        return None
    return f"{d}, {partido}, Buenos Aires, Argentina"

# ---------------------------------------------------------------------------
# 5) Orquestación
# ---------------------------------------------------------------------------
def main():
    do_geocode = "--geocode" in sys.argv
    print("Extrayendo efectores...")
    efectores = extract_efectores()
    print(f"  {len(efectores)} registros (deduplicados)")
    efectores += caps_tres_de_febrero(efectores)
    efectores = unificar(efectores)
    print(f"  {len(efectores)} establecimientos únicos (unificados por partido + nombre)")
    asignar_servicios_pdf(efectores)

    # Municipios
    munis = sorted(set(r['partido'] for r in efectores if r['partido']))
    municipios = []
    for m in munis:
        coord = MUNICIPIO_COORD_MANUAL.get(m) or (geocode(f"{m}, Buenos Aires, Argentina") if do_geocode else None)
        if coord and 'lat' in coord and not dentro_rs7(coord['lat'], coord['lon']):
            print(f"  [AVISO] el centro del municipio {m} cae fuera de la RS VII: revisar MUNICIPIO_COORD_MANUAL")
        municipios.append({
            'nombre': m,
            'lat': coord['lat'] if coord else None,
            'lon': coord['lon'] if coord else None,
            'display_name': coord.get('display_name', '') if coord else '',
        })
    print(f"  {len(municipios)} municipios georreferenciados")

    # Coordenadas por municipio (para fallback)
    mun_coord = {m['nombre']: (m['lat'], m['lon']) for m in municipios}

    # Georreferenciar efectores
    n_geo = 0
    for i, r in enumerate(efectores):
        coord = None
        for dom in r.pop('_doms_geo', [r['domicilio']]):
            q = clean_address(dom, r['partido'])
            if do_geocode and q:
                c = geocode(q)
                if c and 'lat' in c and dentro_rs7(c['lat'], c['lon']):
                    coord = c
                    break
                coord = coord or c
        if coord and 'lat' in coord and not dentro_rs7(coord['lat'], coord['lon']):
            print(f"  [AVISO] coordenada fuera de la RS VII descartada: {r['nombre']} ({r['partido']}) "
                  f"-> {coord['lat']:.4f}, {coord['lon']:.4f}")
            coord = None
        if coord and 'lat' in coord:
            r['lat'] = coord['lat']
            r['lon'] = coord['lon']
            r['geo_origen'] = 'domicilio'
            n_geo += 1
        else:
            # fallback: centroide del municipio + jitter determinístico
            mc = mun_coord.get(r['partido'])
            if mc and mc[0] is not None:
                seed = sum(ord(c) for c in r['nombre'])
                jitter = ((seed % 100) / 10000.0) - 0.005
                r['lat'] = round(mc[0] + jitter, 5)
                r['lon'] = round(mc[1] + jitter * 0.7, 5)
                r['geo_origen'] = 'centroide_municipio'
            else:
                r['lat'] = None
                r['lon'] = None
                r['geo_origen'] = 'sin_geo'
        if i % 50 == 0:
            print(f"    geocodificando {i}/{len(efectores)}...")

    if do_geocode:
        print(f"  georreferenciados por domicilio: {n_geo}/{len(efectores)}")

    # Taxonomía
    tax = build_taxonomia()

    # Resumen
    from collections import Counter
    resumen = {
        'total_efectores': len(efectores),
        'total_municipios': len(municipios),
        'por_municipio': dict(Counter(r['partido'] for r in efectores)),
        'por_categoria': dict(Counter(s['categoria'] for r in efectores for s in r['servicios'])),
        'por_subcategoria': dict(Counter(s['subcategoria'] for r in efectores for s in r['servicios'])),
        'por_nivel': dict(Counter(r['nivel'] if r['nivel'] else 'Sin nivel' for r in efectores)),
        'categorias': len(tax),
    }

    # Escribir salidas
    json.dump(municipios, open(os.path.join(DATA, "municipios.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    json.dump(efectores, open(os.path.join(DATA, "efectores.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    json.dump(tax, open(os.path.join(DATA, "taxonomia.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    json.dump(resumen, open(os.path.join(DATA, "resumen.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    print("\nResumen:")
    print(json.dumps(resumen, ensure_ascii=False, indent=2))
    print("\nSalidas escritas en", DATA)

# Nombres canónicos: la taxonomía debe usar EXACTAMENTE los mismos textos que
# `efectores.json`, porque el panel filtra por igualdad de texto.
CATEGORIA_CANON = {
    "EFECTORES DE SALUD PUBLICA": "EFECTORES DE SALUD PÚBLICA",
    "NUTRICION Y ALIMENTACION": "NUTRICIÓN Y ALIMENTACIÓN",
}
SUBCATEGORIA_CANON = {
    "PRIMER NIVEL": "Primer Nivel",
    "SEGUNDO NIVEL": "Segundo Nivel",
    "TERCER NIVEL": "Tercer Nivel",   # en el Excel tiene una nota pegada ("entrar a poblaciones...")
}

def canon_categoria(c):
    c = str(c).strip()
    return CATEGORIA_CANON.get(_norm(c), c)

def canon_subcategoria(s):
    s = str(s).strip()
    n = _norm(s)
    for k, v in SUBCATEGORIA_CANON.items():
        if n == k or n.startswith(k + " "):
            return v
    return s

def build_taxonomia():
    import openpyxl
    wb = openpyxl.load_workbook(os.path.join(SOURCE, "REPOSITORIO DE ENLACES GEOREF.xlsx"), data_only=True)
    ws = wb['Hoja 1']
    tax, cur = {}, None
    for row in ws.iter_rows(values_only=True):
        c, s, a, m, e, o = row
        if c and c not in ('CATEGORIA',):
            c = canon_categoria(c)
            if c not in tax:
                tax[c] = {}
            cur = c
        if s and s not in ('SUBCATEGORIA',) and cur and s:
            s = canon_subcategoria(s)
            if s not in tax[cur]:
                tax[cur][s] = set()
            if m and m not in ('MUNICIPIO',):
                tax[cur][s].add(canon_mun(m))
    # El nivel de atención (Primer/Segundo/Tercer) NO es un filtro del panel:
    # se distingue por color en el mapa. Se quita esa categoría del menú.
    tax.pop(CATEGORIA_CANON["EFECTORES DE SALUD PUBLICA"], None)
    out = []
    for c, subs in tax.items():
        out.append({'categoria': c, 'subcategorias': [
            {'nombre': s, 'municipios': sorted(munis)} for s, munis in subs.items()
        ]})
    return out

if __name__ == "__main__":
    main()
